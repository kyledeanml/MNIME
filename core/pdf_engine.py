"""
High-Performance PDF Engine implementing core document transformations.
Uses PyMuPDF (C-level rendering) for blazing-fast batch operations with seamless
fallbacks to pypdf and Pillow.
"""

import os
import io
from typing import List, Callable, Optional
from .file_item import FileItem


class PDFEngine:
    """Core backend engine optimized for high-throughput batch document processing."""

    @staticmethod
    def combine_files(
        file_items: List[FileItem],
        output_path: str,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> str:
        """
        Combines multiple PDF files and images in sequential order into a single PDF.
        Optimized with PyMuPDF for lightning-fast multi-file merging.
        """
        total_items = len(file_items)
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 1. Primary High-Speed Engine: PyMuPDF
        try:
            import pymupdf

            merged_doc = pymupdf.open()
            open_docs = []
            for idx, item in enumerate(file_items):
                if progress_callback:
                    pct = int((idx / total_items) * 90)
                    progress_callback(pct, f"Merging {item.file_name} ({idx + 1}/{total_items})...")

                ext = item.extension.lower()
                if ext == ".pdf":
                    sub_doc = pymupdf.open(item.file_path)
                    merged_doc.insert_pdf(sub_doc)
                    open_docs.append(sub_doc)
                elif ext in [".jpg", ".jpeg", ".png", ".bmp", ".webp"]:
                    img_doc = pymupdf.open(item.file_path)
                    pdf_bytes = img_doc.convert_to_pdf()
                    img_pdf = pymupdf.open("pdf", pdf_bytes)
                    merged_doc.insert_pdf(img_pdf)
                    open_docs.append(img_pdf)
                    open_docs.append(img_doc)
                elif ext == ".txt":
                    with open(item.file_path, "r", encoding="utf-8") as f:
                        text_content = f.read()
                    txt_pdf = pymupdf.open()
                    page = txt_pdf.new_page()
                    rect = pymupdf.Rect(50, 50, page.rect.width - 50, page.rect.height - 50)
                    page.insert_textbox(rect, text_content, fontsize=12, fontname="helv")
                    merged_doc.insert_pdf(txt_pdf)
                    open_docs.append(txt_pdf)

            if progress_callback:
                progress_callback(95, "Compacting and saving document...")

            merged_doc.save(output_path, garbage=3, deflate=True)
            merged_doc.close()
            
            for d in open_docs:
                d.close()

            if progress_callback:
                progress_callback(100, f"Finished! Merged {total_items} files.")

            return output_path

        except Exception as e:
            # 2. Fallback Engine: pypdf & Pillow
            import pypdf
            from PIL import Image

            writer = pypdf.PdfWriter()
            for idx, item in enumerate(file_items):
                if progress_callback:
                    pct = int((idx / total_items) * 90)
                    progress_callback(pct, f"Processing {item.file_name} (fallback)...")

                ext = item.extension.lower()
                if ext == ".pdf":
                    reader = pypdf.PdfReader(item.file_path)
                    for page in reader.pages:
                        writer.add_page(page)
                elif ext in [".jpg", ".jpeg", ".png", ".bmp", ".webp"]:
                    with Image.open(item.file_path) as img:
                        img_rgb = img.convert("RGB")
                        temp_pdf_bytes = io.BytesIO()
                        img_rgb.save(temp_pdf_bytes, format="PDF", resolution=150.0)
                        temp_pdf_bytes.seek(0)
                        img_reader = pypdf.PdfReader(temp_pdf_bytes)
                        for page in img_reader.pages:
                            writer.add_page(page)
                elif ext == ".txt":
                    with open(item.file_path, "r", encoding="utf-8") as f:
                        text_content = f.read()
                    
                    from PIL import Image, ImageDraw
                    img = Image.new('RGB', (850, 1100), color=(255, 255, 255))
                    d = ImageDraw.Draw(img)
                    d.text((50, 50), text_content, fill=(0,0,0))
                    
                    temp_pdf_bytes = io.BytesIO()
                    img.save(temp_pdf_bytes, format="PDF", resolution=150.0)
                    temp_pdf_bytes.seek(0)
                    txt_reader = pypdf.PdfReader(temp_pdf_bytes)
                    for page in txt_reader.pages:
                        writer.add_page(page)

            if progress_callback:
                progress_callback(95, "Writing output file...")

            with open(output_path, "wb") as f_out:
                writer.write(f_out)

            if progress_callback:
                progress_callback(100, f"Finished! Merged {total_items} files.")

            return output_path

    @staticmethod
    def convert_jpg_to_pdf(
        image_items: List[FileItem],
        output_path: str,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> str:
        """
        Converts one or multiple images into a single clean PDF document.
        """
        if not image_items:
            raise ValueError("No images provided for conversion.")

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        total = len(image_items)

        # 1. Primary Engine: PyMuPDF
        try:
            import pymupdf
            import concurrent.futures

            def convert_image(item):
                try:
                    img_doc = pymupdf.open(item.file_path)
                    pdf_bytes = img_doc.convert_to_pdf()
                    img_doc.close()
                    return pdf_bytes
                except Exception:
                    return None
            
            pdf_bytes_list = [None] * total
            completed = 0
            
            with concurrent.futures.ThreadPoolExecutor(max_workers=os.cpu_count()) as executor:
                futures = {executor.submit(convert_image, item): idx for idx, item in enumerate(image_items)}
                for future in concurrent.futures.as_completed(futures):
                    idx = futures[future]
                    pdf_bytes_list[idx] = future.result()
                    completed += 1
                    if progress_callback:
                        pct = int((completed / total) * 80)
                        progress_callback(pct, f"Converted {completed}/{total} images using {os.cpu_count()} cores...")

            pdf_doc = pymupdf.open()
            open_docs = []
            
            for b in pdf_bytes_list:
                if b:
                    img_pdf = pymupdf.open("pdf", b)
                    pdf_doc.insert_pdf(img_pdf)
                    open_docs.append(img_pdf)

            if progress_callback:
                progress_callback(95, "Optimizing PDF output...")

            pdf_doc.save(output_path, garbage=3, deflate=True)
            pdf_doc.close()

            for d in open_docs:
                d.close()

            if progress_callback:
                progress_callback(100, "Done!")

            return output_path

        except Exception:
            # 2. Fallback Engine: Pillow
            from PIL import Image

            images = []
            for i, item in enumerate(image_items):
                if progress_callback:
                    progress_callback(int((i / total) * 80), f"Loading {item.file_name}...")
                img = Image.open(item.file_path).convert("RGB")
                images.append(img)

            first_img = images[0]
            remaining = images[1:] if len(images) > 1 else []

            if progress_callback:
                progress_callback(90, "Saving PDF...")

            first_img.save(
                output_path,
                "PDF",
                resolution=150.0,
                save_all=True,
                append_images=remaining,
            )

            for img in images:
                img.close()

            if progress_callback:
                progress_callback(100, "Done!")

            return output_path

    @staticmethod
    def split_pdf(
        pdf_item: FileItem,
        output_dir: str,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> List[str]:
        """
        Splits a PDF into individual 1-page PDF files.
        """
        output_files: List[str] = []
        os.makedirs(output_dir, exist_ok=True)
        base_name = os.path.splitext(pdf_item.file_name)[0]

        try:
            import pymupdf
            import concurrent.futures
            from PyQt6.QtCore import QSettings
            from core.nlp_engine import NLPEngine

            settings = QSettings("MNIME", "MNIMEApp")
            use_nlp = str(settings.value("nlp_smart_indexing", "true")).lower() == "true"
            nlp_engine = NLPEngine.get_instance()
            if use_nlp:
                nlp_engine.check_model()
            is_nlp_active = use_nlp and nlp_engine.is_loaded

            doc = pymupdf.open(pdf_item.file_path)
            total_pages = len(doc)
            doc.close()

            def process_page(page_num):
                try:
                    local_doc = pymupdf.open(pdf_item.file_path)
                    page = local_doc[page_num]
                    fallback_name = f"{base_name}_page_{page_num + 1:03d}"
                    final_name = fallback_name

                    if is_nlp_active:
                        page_text = page.get_text("text").strip()
                        smart_name = nlp_engine.generate_smart_filename(page_text, fallback_name)
                        # Append page number to guarantee uniqueness
                        if smart_name != fallback_name:
                            final_name = f"{smart_name}_p{page_num + 1:03d}"
                            
                    new_doc = pymupdf.open()
                    new_doc.insert_pdf(local_doc, from_page=page_num, to_page=page_num)
                    out_path = os.path.join(output_dir, f"{final_name}.pdf")
                    new_doc.save(out_path, garbage=3, deflate=True)
                    new_doc.close()
                    local_doc.close()
                    return out_path
                except Exception as e:
                    print(f"Error splitting page {page_num}: {e}")
                    return None

            completed = 0
            with concurrent.futures.ThreadPoolExecutor(max_workers=os.cpu_count()) as executor:
                futures = {executor.submit(process_page, i): i for i in range(total_pages)}
                for future in concurrent.futures.as_completed(futures):
                    out_path = future.result()
                    if out_path:
                        output_files.append(out_path)
                    completed += 1
                    if progress_callback:
                        pct = int((completed / total_pages) * 95)
                        progress_callback(pct, f"Split {completed}/{total_pages} pages...")

            output_files.sort()

            if progress_callback:
                progress_callback(100, f"Saved {len(output_files)} pages.")

            return output_files

        except Exception as e:
            raise RuntimeError(f"Error splitting PDF: {e}")

    @staticmethod
    def convert_pdf_to_jpg(
        pdf_item: FileItem,
        output_dir: str,
        dpi: int = 200,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> List[str]:
        """
        Converts pages of a PDF document into individual high-resolution JPG images.
        """
        output_files: List[str] = []
        os.makedirs(output_dir, exist_ok=True)
        base_name = os.path.splitext(pdf_item.file_name)[0]

        try:
            import pymupdf
            import concurrent.futures
            
            # Open just to get page count, then close
            doc = pymupdf.open(pdf_item.file_path)
            total_pages = len(doc)
            doc.close()

            def render_page(page_num):
                try:
                    local_doc = pymupdf.open(pdf_item.file_path)
                    page = local_doc[page_num]
                    zoom = dpi / 72.0
                    matrix = pymupdf.Matrix(zoom, zoom)
                    pix = page.get_pixmap(matrix=matrix, alpha=False)
                    out_path = os.path.join(output_dir, f"{base_name}_page_{page_num + 1:03d}.jpg")
                    pix.save(out_path)
                    local_doc.close()
                    return out_path
                except Exception as e:
                    print(f"Error rendering page {page_num}: {e}")
                    return None

            completed = 0
            with concurrent.futures.ThreadPoolExecutor(max_workers=os.cpu_count()) as executor:
                futures = {executor.submit(render_page, i): i for i in range(total_pages)}
                for future in concurrent.futures.as_completed(futures):
                    out_path = future.result()
                    if out_path:
                        output_files.append(out_path)
                    completed += 1
                    if progress_callback:
                        pct = int((completed / total_pages) * 95)
                        progress_callback(pct, f"Rendered {completed}/{total_pages} pages using {os.cpu_count()} cores...")
            
            output_files.sort()

        except Exception as e:
            raise RuntimeError(f"Error rendering PDF to images: {e}")

        if progress_callback:
            progress_callback(100, f"Saved {len(output_files)} images.")

        return output_files

    @staticmethod
    def compress_pdf(
        pdf_item: FileItem,
        output_path: str,
        compression_level: str = "medium",
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> str:
        """
        Compresses a PDF using PyMuPDF stream deflation, xref compaction,
        and duplicate object elimination.
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        try:
            import pymupdf

            if progress_callback:
                progress_callback(20, "Analyzing PDF structure for compression...")

            doc = pymupdf.open(pdf_item.file_path)

            if progress_callback:
                progress_callback(60, "Deflating streams & compacting objects...")

            doc.save(output_path, garbage=4, deflate=True, clean=True)
            doc.close()

            if progress_callback:
                progress_callback(100, "Compression complete!")

            return output_path

        except Exception:
            # Fallback to pypdf
            import pypdf

            if progress_callback:
                progress_callback(10, "Opening PDF for compression...")

            reader = pypdf.PdfReader(pdf_item.file_path)
            writer = pypdf.PdfWriter()

            total_pages = len(reader.pages)
            for i, page in enumerate(reader.pages):
                if progress_callback:
                    pct = 10 + int((i / total_pages) * 70)
                    progress_callback(pct, f"Compressing page {i + 1} of {total_pages}...")

                page.compress_content_streams()
                writer.add_page(page)

            if progress_callback:
                progress_callback(85, "Optimizing and saving compressed file...")

            with open(output_path, "wb") as f_out:
                writer.write(f_out)

            if progress_callback:
                progress_callback(100, "Compression complete!")

            return output_path

    @staticmethod
    def convert_pdf_to_docx(
        pdf_item: FileItem,
        output_path: str,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> str:
        """
        Converts a PDF file to a Microsoft Word (.docx) document.
        """
        try:
            from pdf2docx import Converter

            if progress_callback:
                progress_callback(10, "Initializing PDF to DOCX converter...")

            cv = Converter(pdf_item.file_path)
            if progress_callback:
                progress_callback(40, "Parsing layout and extracting content...")

            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
            cv.convert(output_path, start=0, end=None)
            cv.close()

            if progress_callback:
                progress_callback(100, "Word document generated successfully!")

            return output_path

        except ImportError:
            try:
                import pypdf
                from docx import Document

                if progress_callback:
                    progress_callback(20, "Extracting text with pypdf...")

                reader = pypdf.PdfReader(pdf_item.file_path)
                doc = Document()
                doc.add_heading(os.path.splitext(pdf_item.file_name)[0], 0)

                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        doc.add_paragraph(text)
                    doc.add_page_break()

                doc.save(output_path)
                if progress_callback:
                    progress_callback(100, "Document saved (text extracted).")
                return output_path

            except Exception as e:
                raise RuntimeError(
                    f"Please install pdf2docx (`pip install pdf2docx`) for full formatting fidelity: {e}"
                )

    @staticmethod
    def bookmark(
        pdf_item: FileItem,
        output_path: str,
        progress_callback: Optional[Callable[[int, str], None]] = None,
    ) -> str:
        """
        Automatically adds bookmarks to a PDF by extracting the most prominent text
        (e.g., largest font size) from each page to use as a section header.
        """
        try:
            import pymupdf
            from PyQt6.QtCore import QSettings
            from core.nlp_engine import NLPEngine
            
            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

            settings = QSettings("MNIME", "MNIMEApp")
            use_nlp = str(settings.value("nlp_smart_indexing", "true")).lower() == "true"
            nlp_engine = NLPEngine.get_instance()
            if use_nlp:
                nlp_engine.check_model()
            is_nlp_active = use_nlp and nlp_engine.is_loaded

            if progress_callback:
                progress_callback(5, "Opening PDF for analysis...")

            doc = pymupdf.open(pdf_item.file_path)
            toc = []

            total_pages = len(doc)
            
            # Step 1: Sample the first few pages to find the most common (base) font size
            import re
            from collections import Counter
            font_sizes = Counter()
            sample_pages = min(10, total_pages)
            for i in range(sample_pages):
                blocks = doc[i].get_text("dict").get("blocks", [])
                for b in blocks:
                    if b.get("type") == 0:
                        for l in b.get("lines", []):
                            for s in l.get("spans", []):
                                text = s.get("text", "").strip()
                                size = round(s.get("size", 0), 1)
                                if text and size > 0:
                                    font_sizes[size] += len(text)
                                    
            base_size = 10.0
            if font_sizes:
                # The font size with the most characters is likely the body text
                base_size = font_sizes.most_common(1)[0][0]

            # Regex for explicit chapter/section markers
            chapter_pattern = re.compile(r"^(chapter|part|section|appendix|unit|module)\s+([a-z0-9\.\-]+)", re.IGNORECASE)

            all_largest_texts = []

            for i in range(total_pages):
                if progress_callback:
                    pct = 10 + int((i / total_pages) * 70)
                    progress_callback(pct, f"Analyzing page {i+1}/{total_pages}...")

                page = doc[i]
                blocks = page.get_text("dict").get("blocks", [])

                largest_size = -1
                best_text = ""
                has_explicit_chapter = False
                page_largest_text = f"Page {i+1}"
                page_absolute_largest_size = -1

                for b in blocks:
                    if b.get("type") == 0:  # Text block
                        for l in b.get("lines", []):
                            for s in l.get("spans", []):
                                text = s.get("text", "").strip()
                                size = s.get("size", 0)
                                flags = s.get("flags", 0)
                                is_bold = bool(flags & 16) # Bit 4 is bold in PyMuPDF
                                
                                if not text or len(text) > 100:
                                    continue
                                    
                                # Track absolute largest text for fallback
                                if size > page_absolute_largest_size:
                                    page_absolute_largest_size = size
                                    page_largest_text = text
                                    
                                # Rule 1: Explicit match
                                if chapter_pattern.match(text):
                                    best_text = text
                                    has_explicit_chapter = True
                                    break
                                    
                                # Rule 2: Largest text on page, and significantly larger than body text
                                # Or slightly larger and Bold
                                if size > largest_size and (size >= (base_size * 1.15) or (is_bold and size >= base_size * 1.05)):
                                    largest_size = size
                                    best_text = text
                            
                            if has_explicit_chapter:
                                break
                    if has_explicit_chapter:
                        break

                all_largest_texts.append(page_largest_text)

                # Only add a bookmark if we confidently found a heading
                if best_text:
                    if is_nlp_active:
                        if progress_callback:
                            progress_callback(pct, f"NLP summarizing page {i+1}/{total_pages}...")
                        page_text = page.get_text("text").strip()
                        best_text = nlp_engine.generate_verbose_bookmark(best_text, page_text)
                        
                    toc.append([1, best_text, i + 1])
                    
            if not toc:
                # Fallback: if heuristics failed completely (e.g. plain text or images only),
                # just bookmark every page using the largest text found
                for i, text in enumerate(all_largest_texts):
                    toc.append([1, text, i + 1])
                
                # Explicitly yield GIL to prevent main thread cursor animations from lagging
                import time
                time.sleep(0.005)

            if progress_callback:
                progress_callback(85, "Generating Table of Contents...")

            doc.set_toc(toc)

            if progress_callback:
                progress_callback(90, "Saving bookmarked PDF...")

            doc.save(output_path, garbage=3, deflate=True)
            doc.close()

            if progress_callback:
                progress_callback(100, "Done! Bookmarks added.")

            return output_path

        except Exception as e:
            raise RuntimeError(f"Error bookmarking PDF: {e}")
